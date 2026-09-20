# -*- coding: utf-8 -*-
# SPDX-License-Identifier: MIT
# Copyright (c) 2026 Damian Patricio Soto
"""
Reporte 1: Reporte Excel (.xlsx).

Consolida en un único libro, en hojas separadas, la información de los reportes
tabulares del expediente:
  · Resumen            (datos de carátula)
  · Índice             (un registro por documento)
  · Firmantes          (un registro por firmante)
  · Destinatarios      (ME/NO)
  · Listado Embebidos
  · Orden documentos   (verificación de orden cronológico)
"""
from typing import List, Optional

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_OK = True
except ImportError:
    OPENPYXL_OK = False

from ..config import ruta_reporte
from ..core.models import Documento
from .listado_embebidos import _formatear_tamaño
from .caratula_y_orden import (
    _firma_mas_antigua,
    _firma_mas_reciente,
    _calcular_dias_caratula,
    _calcular_dias_firmas,
    _verificar_orden,
)

_INDICE_HEADERS = [
    "N° Orden", "N° Documento", "Fecha", "Referencia", "Páginas",
    "Embebidos", "Firmantes", "Último Firmante", "Cargo", "Área",
]
_FIRMANTES_HEADERS = [
    "N° Orden", "Código Documento", "Fecha Documento",
    "Firmante", "Fecha Firma", "Cargo", "Área",
]
_DESTINATARIOS_HEADERS = [
    "N° Orden", "Código Documento", "Fecha Documento",
    "A / Copia A", "Nombre", "Área",
]
_EMBEBIDOS_HEADERS = [
    "N° Orden", "N° Documento", "Fecha Documento",
    "Archivo Embebido", "Tamaño",
]


def generar_excel(
    doc_caratula: Optional[Documento],
    documentos: List[Documento],
    output_dir: str,
) -> str:
    """Genera 'Reporte <carpeta>.xlsx' en output_dir. Devuelve la ruta del archivo generado."""
    if not OPENPYXL_OK:
        raise RuntimeError("openpyxl no está instalado. No se puede generar el Reporte Excel.")

    wb = Workbook()

    ws_resumen = wb.active
    ws_resumen.title = "Resumen"
    _hoja_resumen(ws_resumen, doc_caratula, documentos)

    _hoja_tabla(wb.create_sheet("Índice"), _INDICE_HEADERS, _filas_indice(documentos))
    _hoja_tabla(wb.create_sheet("Firmantes"), _FIRMANTES_HEADERS, _filas_firmantes(documentos))
    _hoja_tabla(wb.create_sheet("Destinatarios"), _DESTINATARIOS_HEADERS, _filas_destinatarios(documentos))
    _hoja_tabla(wb.create_sheet("Listado Embebidos"), _EMBEBIDOS_HEADERS, _filas_embebidos(documentos))
    _hoja_orden_documentos(wb.create_sheet("Orden documentos"), documentos)

    salida = ruta_reporte(output_dir, "Reporte", ".xlsx")
    wb.save(salida)
    return str(salida)


# ─── Construcción de filas (comparte lógica con los reportes CSV) ──────────────

def _filas_indice(documentos: List[Documento]) -> List[list]:
    filas = []
    for doc in documentos:
        uf = doc.ultimo_firmante
        filas.append([
            doc.numero_orden, doc.codigo, doc.fecha_documento, doc.referencia,
            doc.cantidad_paginas, doc.cantidad_embebidos, doc.cantidad_firmantes,
            uf.nombre.upper() if uf else "", uf.cargo if uf else "", uf.area if uf else "",
        ])
    return filas


def _filas_firmantes(documentos: List[Documento]) -> List[list]:
    filas = []
    for doc in documentos:
        for f in doc.firmantes:
            filas.append([
                doc.numero_orden, doc.codigo, doc.fecha_documento,
                f.nombre.upper(), f.fecha_firma, f.cargo, f.area,
            ])
    return filas


def _filas_destinatarios(documentos: List[Documento]) -> List[list]:
    filas = []
    for doc in documentos:
        if doc.tipo not in ("ME", "NO"):
            continue
        for d in doc.destinatarios:
            filas.append([
                doc.numero_orden, doc.codigo, doc.fecha_documento,
                d.tipo, d.nombre, d.area,
            ])
    return filas


def _filas_embebidos(documentos: List[Documento]) -> List[list]:
    filas = []
    for doc in documentos:
        if doc.cantidad_embebidos == 0:
            continue
        for e in doc.embebidos:
            filas.append([
                doc.numero_orden, doc.codigo, doc.fecha_documento,
                e.nombre, _formatear_tamaño(len(e.data)),
            ])
    return filas


# ─── Construcción de hojas ─────────────────────────────────────────────────────

def _estilos():
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    return header_fill, header_font, border


def _hoja_tabla(ws, headers: List[str], filas: List[list]):
    header_fill, header_font, border = _estilos()

    ws.append(headers)
    for col, _ in enumerate(headers, 1):
        celda = ws.cell(row=1, column=col)
        celda.fill = header_fill
        celda.font = header_font
        celda.alignment = Alignment(horizontal="center", vertical="center")
        celda.border = border

    for fila in filas:
        ws.append(fila)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(filas) + 1)}"
    _ajustar_anchos(ws, headers, filas)


def _ajustar_anchos(ws, headers: List[str], filas: List[list]):
    for col, header in enumerate(headers, 1):
        ancho = len(str(header))
        for fila in filas:
            valor = fila[col - 1] if col - 1 < len(fila) else ""
            ancho = max(ancho, len(str(valor)))
        ws.column_dimensions[get_column_letter(col)].width = min(ancho + 2, 70)


def _hoja_resumen(ws, doc_caratula: Optional[Documento], documentos: List[Documento]):
    total_docs      = len(documentos)
    total_embebidos = sum(d.cantidad_embebidos for d in documentos)
    total_hojas     = sum(d.cantidad_paginas for d in documentos)
    firma_ant       = _firma_mas_antigua(documentos)
    firma_rec       = _firma_mas_reciente(documentos)
    fecha_caratula  = doc_caratula.fecha_documento if doc_caratula else ""
    dias_caratula   = _calcular_dias_caratula(fecha_caratula, firma_rec)
    dias_firmas     = _calcular_dias_firmas(firma_ant, firma_rec)

    def _firma_str(f):
        return f"{f['fecha']}  (doc. {f['orden']})" if f else "(sin firmas detectadas)"

    titulo = ws.cell(row=1, column=1, value="RESUMEN DEL EXPEDIENTE")
    titulo.font = Font(bold=True, size=14, color="1F4E78")

    referencia = doc_caratula.referencia if doc_caratula else ""
    c_ref = ws.cell(row=3, column=1, value=referencia)
    c_ref.font = Font(bold=True)

    filas = [
        ("Número de orden",   doc_caratula.numero_orden if doc_caratula else ""),
        ("", ""),
        ("Total de documentos",             total_docs),
        ("Total de archivos embebidos",     total_embebidos),
        ("Total de hojas",                  total_hojas),
        ("Firma más antigua",               _firma_str(firma_ant)),
        ("Firma más reciente",              _firma_str(firma_rec)),
        ("Días entre firma más antigua y más reciente", dias_firmas),
        ("Días entre carátula y firma más reciente",    dias_caratula),
    ]

    fila_inicio = 4
    for i, (etiqueta, valor) in enumerate(filas):
        r = fila_inicio + i
        c_et = ws.cell(row=r, column=1, value=etiqueta)
        ws.cell(row=r, column=2, value=valor)
        if etiqueta:
            c_et.font = Font(bold=True)

    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 45


def _hoja_orden_documentos(ws, documentos: List[Documento]):
    """Hoja final: verificación de orden cronológico y, si aplica, orden sugerido."""
    header_fill, header_font, border = _estilos()
    orden = _verificar_orden(documentos)

    c_et = ws.cell(row=1, column=1, value="Documentos en orden cronológico")
    c_et.font = Font(bold=True)
    ws.cell(row=1, column=2, value="Sí" if orden["ordenado"] else "No")

    if not orden["ordenado"] and orden["orden_sugerido"]:
        r = 3
        ws.cell(row=r, column=1, value="ORDEN POR FECHA DE ÚLTIMO FIRMANTE").font = Font(
            bold=True, color="1F4E78"
        )
        r += 1
        for col, h in enumerate(["N° Orden", "Fecha firma último firmante"], 1):
            celda = ws.cell(row=r, column=col, value=h)
            celda.fill = header_fill
            celda.font = header_font
            celda.alignment = Alignment(horizontal="center")
            celda.border = border
        for item in orden["orden_sugerido"]:
            r += 1
            ws.cell(row=r, column=1, value=item["orden"])
            ws.cell(row=r, column=2, value=item["fecha_firma"])

    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 45
